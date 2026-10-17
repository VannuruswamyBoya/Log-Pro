import io
import csv
from datetime import datetime, timedelta
from rest_framework.views import APIView
from rest_framework.response import Response
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from core.db import get_submissions_col, get_users_col, log_audit_event
from core.permissions import IsMongoAdmin
from core.deadline_policy import get_current_business_date
from core.views_auth import get_client_ip

def get_report_date_range(request):
    """
    Computes start_date and end_date. Defaults to the last 7 days including today.
    """
    end_str = request.GET.get('end_date')
    start_str = request.GET.get('start_date')

    now = datetime.now()
    if not end_str:
        end_str = get_current_business_date(now)
    if not start_str:
        start_dt = now - timedelta(days=6)
        start_str = start_dt.strftime('%Y-%m-%d')

    return start_str, end_str

def fetch_report_data(start_date, end_date):
    submissions_col = get_submissions_col()
    users_col = get_users_col()

    users = list(users_col.find({'role': 'user', 'is_active': True}))
    user_map = {str(u['_id']): u for u in users}

    query = {
        'is_active': True,
        'business_date': {'$gte': start_date, '$lte': end_date}
    }
    subs = list(submissions_col.find(query).sort('business_date', -1))

    total_submitted = len(subs)
    on_time = sum(1 for s in subs if s.get('status') == 'ON_TIME')
    late = sum(1 for s in subs if s.get('status') == 'LATE')

    # Calculate days in interval
    start_dt = datetime.strptime(start_date, '%Y-%m-%d')
    end_dt = datetime.strptime(end_date, '%Y-%m-%d')
    days_count = max(1, (end_dt - start_dt).days + 1)
    expected_submissions = len(users) * days_count
    missing = max(0, expected_submissions - total_submitted)

    compliance_rate = round((on_time / expected_submissions * 100), 1) if expected_submissions > 0 else 0.0

    return {
        'start_date': start_date,
        'end_date': end_date,
        'total_users': len(users),
        'expected_submissions': expected_submissions,
        'total_submitted': total_submitted,
        'on_time': on_time,
        'late': late,
        'missing': missing,
        'compliance_rate': compliance_rate,
        'submissions': subs,
        'users': users
    }

class WeeklyReportSummaryView(APIView):
    """Returns JSON summary data for the weekly/periodic report dashboard."""
    permission_classes = [IsMongoAdmin]

    def get(self, request):
        start_date, end_date = get_report_date_range(request)
        data = fetch_report_data(start_date, end_date)

        # Build user breakdown
        user_summary = []
        for u in data['users']:
            uid = str(u['_id'])
            user_subs = [s for s in data['submissions'] if s.get('user_id') == uid]
            u_on_time = sum(1 for s in user_subs if s.get('status') == 'ON_TIME')
            u_late = sum(1 for s in user_subs if s.get('status') == 'LATE')
            user_summary.append({
                'user_id': uid,
                'name': u.get('full_name'),
                'email': u.get('email'),
                'department': u.get('department'),
                'submitted': len(user_subs),
                'on_time': u_on_time,
                'late': u_late
            })

        return Response({
            'period': {'start_date': start_date, 'end_date': end_date},
            'kpis': {
                'total_users': data['total_users'],
                'expected_submissions': data['expected_submissions'],
                'total_submitted': data['total_submitted'],
                'on_time': data['on_time'],
                'late': data['late'],
                'missing': data['missing'],
                'compliance_rate': data['compliance_rate']
            },
            'user_breakdown': user_summary,
            'recent_submissions_count': len(data['submissions'])
        })

class ExportReportCSVView(APIView):
    """Exports submissions in the specified period as a standard CSV file."""
    permission_classes = [IsMongoAdmin]

    def get(self, request):
        start_date, end_date = get_report_date_range(request)
        data = fetch_report_data(start_date, end_date)

        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow([
            'Submission ID', 'Business Date', 'Full Name', 'Email', 
            'Department', 'Status', 'Upload Timestamp', 'Original File Name', 
            'Format', 'Size (Bytes)', 'Version', 'Notes'
        ])

        for s in data['submissions']:
            writer.writerow([
                str(s['_id']),
                s.get('business_date'),
                s.get('user_name'),
                s.get('user_email'),
                s.get('department'),
                s.get('status'),
                s.get('upload_timestamp'),
                s.get('original_filename'),
                s.get('file_format'),
                s.get('file_size'),
                s.get('version', 1),
                s.get('notes', '')
            ])

        log_audit_event(request.user.id, request.user.email, 'EXPORT_CSV', {'start_date': start_date, 'end_date': end_date}, get_client_ip(request))

        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="LogPro_Report_{start_date}_to_{end_date}.csv"'
        return response

class ExportReportExcelView(APIView):
    """Exports styled Excel (.xlsx) compliance workbook."""
    permission_classes = [IsMongoAdmin]

    def get(self, request):
        start_date, end_date = get_report_date_range(request)
        data = fetch_report_data(start_date, end_date)

        wb = Workbook()
        ws = wb.active
        ws.title = "Submission Compliance Report"

        # Styling definitions
        navy_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=10)
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        orange_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title Block
        ws.append(["Log Pro - Enterprise Log Sheet Compliance Report"])
        ws['A1'].font = title_font
        ws.append([f"Report Period: {start_date} to {end_date} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
        ws.append([])

        # KPI Summary Block
        ws.append(["KPI Metrics", "Value"])
        ws['A4'].font = bold_font
        ws['B4'].font = bold_font
        kpis = [
            ("Active Users", data['total_users']),
            ("Expected Submissions", data['expected_submissions']),
            ("Total Submissions Received", data['total_submitted']),
            ("On-Time Submissions", data['on_time']),
            ("Late Submissions", data['late']),
            ("Missing Submissions", data['missing']),
            ("Overall Compliance Rate", f"{data['compliance_rate']}%")
        ]
        for label, val in kpis:
            ws.append([label, val])
        ws.append([])

        # Table Header
        headers = [
            "Business Date", "Employee Name", "Email", "Department",
            "Status", "Upload Timestamp", "Filename", "Format", "Size (KB)", "Version"
        ]
        start_row = ws.max_row + 1
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=start_row, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Table Rows
        for s in data['submissions']:
            status_val = s.get('status', '')
            size_kb = round(s.get('file_size', 0) / 1024, 1)
            row_data = [
                s.get('business_date'),
                s.get('user_name'),
                s.get('user_email'),
                s.get('department'),
                status_val,
                s.get('upload_timestamp', '')[:19].replace('T', ' '),
                s.get('original_filename'),
                s.get('file_format'),
                size_kb,
                s.get('version', 1)
            ]
            ws.append(row_data)
            curr_row = ws.max_row
            # Format status cell
            status_cell = ws.cell(row=curr_row, column=5)
            if status_val == 'ON_TIME':
                status_cell.fill = green_fill
                status_cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
            elif status_val == 'LATE':
                status_cell.fill = orange_fill
                status_cell.font = Font(name="Calibri", size=10, bold=True, color="9A3412")

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.font = regular_font

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        log_audit_event(request.user.id, request.user.email, 'EXPORT_EXCEL', {'start_date': start_date, 'end_date': end_date}, get_client_ip(request))

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="LogPro_Compliance_Report_{start_date}_to_{end_date}.xlsx"'
        return response

class ExportReportPDFView(APIView):
    """Generates an executive PDF report formatted cleanly with ReportLab."""
    permission_classes = [IsMongoAdmin]

    def get(self, request):
        start_date, end_date = get_report_date_range(request)
        data = fetch_report_data(start_date, end_date)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=14
        )

        elements = []
        elements.append(Paragraph("Log Pro - Weekly Submission Compliance Report", title_style))
        elements.append(Paragraph(f"Reporting Window: {start_date} to {end_date} &bull; Generated on {datetime.now().strftime('%B %d, %Y at %H:%M')}", subtitle_style))

        # KPI Summary Table
        kpi_data = [
            ['Total Users', 'Expected', 'Received', 'On-Time', 'Late', 'Missing', 'Compliance Rate'],
            [
                str(data['total_users']),
                str(data['expected_submissions']),
                str(data['total_submitted']),
                str(data['on_time']),
                str(data['late']),
                str(data['missing']),
                f"{data['compliance_rate']}%"
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[100, 100, 100, 100, 100, 100, 120])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#F1F5F9')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 11),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 18))

        # Detailed Submissions Table
        elements.append(Paragraph("Submission Log Details", ParagraphStyle('SectionHeading', parent=styles['Heading2'], fontSize=13, textColor=colors.HexColor('#0F172A'), spaceAfter=8)))

        table_rows = [['Date', 'User Name', 'Department', 'Status', 'File Name', 'Format', 'Timestamp']]
        for s in data['submissions'][:35]:  # limit for PDF layout
            table_rows.append([
                s.get('business_date', ''),
                s.get('user_name', '')[:20],
                s.get('department', '')[:15],
                s.get('status', ''),
                s.get('original_filename', '')[:25],
                s.get('file_format', ''),
                s.get('upload_timestamp', '')[:16].replace('T', ' ')
            ])

        if len(table_rows) == 1:
            table_rows.append(['No submissions found for the selected period.', '', '', '', '', '', ''])

        sub_table = Table(table_rows, colWidths=[70, 120, 95, 75, 170, 55, 135])
        sub_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))

        elements.append(sub_table)

        doc.build(elements)
        buffer.seek(0)

        log_audit_event(request.user.id, request.user.email, 'EXPORT_PDF', {'start_date': start_date, 'end_date': end_date}, get_client_ip(request))

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="LogPro_Report_{start_date}_to_{end_date}.pdf"'
        return response
